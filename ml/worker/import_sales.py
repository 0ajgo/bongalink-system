"""Import the hand-kept sales sheet into a FRESH copy of the Bongalink database.

    python import_sales.py --dry-run                  # clean + report, writes nothing
    python import_sales.py                            # (re)creates bongalink_real and imports
    python import_sales.py --file other.xlsx --database bongalink_real2

What it does (see the report it prints):
  * reads tab "SALES" (headers in row 1, data from row 3), skipping blank separator rows;
  * drops services (repairs, resets, labor) and lists them for review;
  * collapses item / model spelling variants with the explicit maps below. An item that is
    in neither ITEM_MAP nor SERVICE_PATTERN stops the run, so nothing is guessed silently;
  * keeps only PAID rows as sales (= orders with status 'completed');
  * creates each product once and one order + order_item per sale, dated with the sheet's date.

Destructive only for the target database: it is dropped and rebuilt from
Backend/database/bongalink_schema.sql + migrations/003_ml_sync.sql every run.
It refuses to touch the production database name ("bongalink").

KNOWN LIMITATIONS (also printed):
  * stock_quantity is NOT real: the sheet has sales, not stock counts. Every product gets
    DEFAULT_STOCK so the dashboard works; recount before trusting reorder suggestions.
  * the sheet has no customers: all orders belong to one placeholder user (import_walkin).
  * orders.payment_method only allows GCASH / BANK_TRANSFER, so Cash, COD, Shopee, LBC and
    mixed payments are stored as BANK_TRANSFER (counted in the report).
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pymysql

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bongalink_ml import config, sqlscript  # noqa: E402

WORKER = Path(__file__).resolve().parent
ROOT = WORKER.parents[1]
SCHEMA = ROOT / "Backend" / "database" / "bongalink_schema.sql"
MIGRATION = ROOT / "Backend" / "database" / "migrations" / "003_ml_sync.sql"
DEFAULT_FILE = WORKER / "data" / "2026 Bongalink Sheet SALES.xlsx"
PRODUCTION_DB = "bongalink"
DEFAULT_STOCK = 5
IMPORT_USER = "import_walkin"

COLUMNS = ["date", "item", "model", "amount", "qty_total", "platform", "payment_method", "payment_status",
           "delivery_method", "delivery_status", "remarks", "date_origin", "service_status"]

# --------------------------------------------------------------------------- cleaning maps
# Labor / service, not stock. Matched against the lower-cased Item text.
SERVICE_PATTERN = re.compile(r"repair|reset|labor|replace|install|overhaul|cleaning|service")

# lower-cased, whitespace-collapsed raw Item -> (clean item, category)
ITEM_MAP = {
    "printhead": ("Printhead", "Printer Parts"),
    "bnew printhead": ("Printhead", "Printer Parts"),          # "Bnew" = brand new, same part
    "head": ("Printhead", "Printer Parts"),                    # only row: Brother T420 head
    "printhead and purge assembly": ("Printhead and Purge Assembly", "Printer Parts"),   # bundle, kept whole
    "mbox": ("Maintenance Box", "Printer Parts"),
    "board": ("Mainboard", "Printer Parts"),
    "bnew board": ("Mainboard", "Printer Parts"),
    "headboard": ("Headboard", "Printer Parts"),
    "cr motor": ("CR Motor", "Printer Parts"),
    "encoder strip": ("Encoder Strip", "Printer Parts"),
    "feeder support gear": ("Feeder Support Gear", "Printer Parts"),
    "control panel flex": ("Control Panel Flex", "Printer Parts"),
    "panel flex bnew": ("Control Panel Flex", "Printer Parts"),
    "flex cable": ("Flex Cable", "Printer Parts"),
    "flex set": ("Flex Set", "Printer Parts"),
    "panel assembly": ("Panel Assembly", "Printer Parts"),
    "paper tray": ("Paper Tray", "Printer Parts"),
    "pickup assembly": ("Pickup Assembly", "Printer Parts"),
    "scanner assembly": ("Scanner Assembly", "Printer Parts"),
    "sensor": ("Sensor", "Printer Parts"),
    "sensor and belt": ("Sensor and Belt", "Printer Parts"),          # bundle, kept whole
    "carriage": ("Carriage", "Printer Parts"),
    "power supply": ("Power Supply", "Printer Parts"),
    "ink": ("Ink Bottle", "Printer Parts"),
    "4 bottles ink": ("Ink Bottle", "Printer Parts"),                 # quantity 4, see QTY_IN_ITEM
    "cartridge": ("Ink Cartridge", "Printer Parts"),
    "ink cartridge": ("Ink Cartridge", "Printer Parts"),
    "cartridge head": ("Cartridge Head", "Printer Parts"),
    "cartridge head set (vat) + j&t": ("Cartridge Head", "Printer Parts"),
    "printer": ("Printer", "Printers"),
    "secondhand printer": ("Printer (Secondhand)", "Printers"),
    "printer and pickup roller": ("Printer and Pickup Roller", "Printers"),   # bundle, kept whole
    "hard drive": ("Hard Drive", "Computer Parts"),
    "keyboard": ("Keyboard", "Computer Parts"),
    "charger": ("Laptop Charger", "Laptop Parts"),
    "a4tech mouse": ("Mouse", "Accessories"),
    "lan cable": ("LAN Cable", "Accessories"),
    "power cord": ("Power Cord", "Accessories"),
    "cable (usb and power cord)": ("USB and Power Cord Cable", "Accessories"),
}
QTY_IN_ITEM = re.compile(r"^(\d+)\s+bottles?\b")
BUNDLES = {"Printhead and Purge Assembly", "Sensor and Belt", "Printer and Pickup Roller", "USB and Power Cord Cable"}

# Maintenance-box part numbers, so a row with no Item but one of these models is still an Mbox.
MBOX_MODELS = {"C9344", "C9345", "T04D1"}

BRAND_WORDS = re.compile(r"^(epson maintenance box|epson|brother|canon|hp|mbox|a4tech)\s+", re.I)


def clean_model(raw, item: str) -> str | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        n = int(raw)
        if item == "Ink Bottle":
            return f"{abs(n):03d}"                       # -3 = "003" mangled by Excel
        return f"L{n}"                                   # 3210 -> L3210 (Epson L-series)
    s = re.sub(r"\s+", " ", str(raw)).strip()
    if s in ("", "--") or s.lower() == item.lower():
        return None
    s = BRAND_WORDS.sub("", s).strip()
    s = BRAND_WORDS.sub("", s).strip()                   # "Mbox T04D1", "Epson Maintenance box C9344"
    up = s.upper()
    if item == "Ink Bottle" and re.fullmatch(r"0-?0-?3( BLANK|INK)?", up):
        return "003"
    up = re.sub(r"^(C934[45])\s*-\s*\S+$", r"\1", up)   # "C9345 - L8050" -> C9345
    up = re.sub(r"^(\d{3,5})$", r"L\1", up)              # "3210" -> L3210
    up = re.sub(r"^L(X|Q)-?(\d+)$", r"L\1-\2", up)       # LX310 / LX-310 -> LX-310, LQ310 -> LQ-310
    up = re.sub(r"^(TM-)?U220$", "TM-U220", up)
    return up


def brand_of(raw_model, clean: str | None) -> str | None:
    text = f"{raw_model or ''} {clean or ''}".lower()
    if "brother" in text or (clean and re.match(r"^T\d", clean) and clean not in MBOX_MODELS):
        return "Brother"
    if "canon" in text or (clean and re.match(r"^G\d", clean)):
        return "Canon"
    if text.strip().startswith("hp") or (clean and clean.startswith("HP")) or (clean and clean.startswith("TRICOLOR HP")):
        return "HP"
    if "acer" in text:
        return "Acer"
    if "a4tech" in text:
        return "A4Tech"
    if clean and (re.match(r"^(L|LX-|LQ-|TM-|M\d|C934|T04D|003)", clean)):
        return "Epson"
    return None


# --------------------------------------------------------------------------- money & quantity
MONEY = re.compile(r"(\d[\d,]*(?:\.\d+)?)\s*(k)?\b", re.I)


def money(text) -> float | None:
    """First money value in a cell: 'P1,600' -> 1600, 'P5.5k' -> 5500, 2550 -> 2550."""
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return float(text) if text > 0 else None
    m = MONEY.search(str(text).replace("P,", "P").replace(", ", ","))
    if not m:
        return None
    v = float(m.group(1).replace(",", ""))
    return v * 1000 if m.group(2) else v


def explicit_qty(text) -> int | None:
    """Quantity written in the cell, or None if the cell doesn't state one."""
    if text is None or isinstance(text, (int, float)):
        return None
    s = str(text)
    m = re.search(r"(\d+)\s*pcs?\b", s, re.I)                    # "1pc - P1,600", "= 30pcs.", "(2pcs.)"
    if m:
        return int(m.group(1))
    m = re.search(r"[xX]\s*(\d+)\b", s)                             # "P2,600 x 3", "P800x10 = 8,000"
    if m:
        return int(m.group(1))
    return None


# --------------------------------------------------------------------------- parsing
@dataclass
class Row:
    sheet_row: int
    raw: dict
    date: dt.date | None = None
    item: str | None = None
    category: str | None = None
    model: str | None = None
    brand: str | None = None
    qty: int = 1
    unit_price: float | None = None
    paid: bool = False
    flags: list[str] = field(default_factory=list)

    @property
    def key(self):
        return (self.item, self.model)


def is_paid(status) -> bool:
    return isinstance(status, str) and re.match(r"^(fully\s+)?paid\b", status.strip(), re.I) is not None


def read_sheet(path: Path):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["SALES"]
    header = [str(c).strip() if c else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    expected = ["Date", "Item", "Model", "Amount", "Quantity x Total", "Platform", "Payment Method", "Payment Status",
                "Delivery Method", "Delivery Status", "Remarks", "Date Origin", "Service Status"]
    if header[:13] != expected:
        raise SystemExit(f"unexpected header row: {header[:13]}")
    for n, values in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        touched = any(v is not None for v in values[:13])     # even a lone space counts as "used"
        cells = dict(zip(COLUMNS, values[:13]))
        cells = {k: (v.strip() or None) if isinstance(v, str) else v for k, v in cells.items()}
        yield n, cells, touched


def clean(path: Path):
    report = {"separator": [], "no_item": [], "services": [], "unpaid": [], "review": [],
              "item_map": defaultdict(Counter), "model_map": defaultdict(Counter), "payment_fallback": Counter()}
    rows: list[Row] = []
    last_date = None

    for n, c, touched in read_sheet(path):
        if not touched:
            continue                                   # never-used rows (the sheet's range runs to 1019)
        has_content = any(c[k] not in (None, "") for k in COLUMNS if k not in ("date", "date_origin"))
        date = c["date"] if isinstance(c["date"], dt.datetime) else None
        origin = c["date_origin"] if isinstance(c["date_origin"], dt.datetime) else None
        if not has_content:
            report["separator"].append(n)
            if date:
                last_date = date.date()
            continue

        r = Row(n, c)
        if date:
            r.date = date.date()
        elif origin:
            r.date = origin.date()
            r.flags.append("date from Date Origin")
        elif last_date:
            r.date = last_date
            r.flags.append(f"no Date or Date Origin: used previous row's date {last_date}")
        last_date = r.date or last_date

        raw_item = re.sub(r"\s+", " ", c["item"] or "").strip()
        key = raw_item.lower()
        if not raw_item:
            model = (c["model"] or "").strip().upper()
            if model in MBOX_MODELS:
                key = "mbox"
                r.flags.append(f"Item blank: inferred Maintenance Box from model {model}")
            else:
                report["no_item"].append((n, c))
                continue
        if SERVICE_PATTERN.search(key):
            report["services"].append((n, c))
            continue
        if key not in ITEM_MAP:
            raise SystemExit(f"row {n}: item {raw_item!r} is not in ITEM_MAP or SERVICE_PATTERN. "
                             "Add it to one of them in import_sales.py.")
        r.item, r.category = ITEM_MAP[key]
        report["item_map"][r.item][raw_item or "(blank)"] += 1
        if r.item in BUNDLES:
            r.flags.append(f"bundle item kept as one product: {raw_item!r}")
        if key in ("panel flex bnew", "head"):
            r.flags.append(f"merged {raw_item!r} into {r.item!r}")

        r.model = clean_model(c["model"], r.item)
        r.brand = brand_of(c["model"], r.model)
        report["model_map"][r.item][f"{c['model']!r} -> {r.model}"] += 1

        # Quantity: "Quantity x Total" first, then Amount, then the item text, else 1.
        q = explicit_qty(c["qty_total"])
        q_src = "Quantity x Total"
        if q is None:
            q, q_src = explicit_qty(c["amount"]), "Amount"
        if q is None and QTY_IN_ITEM.match(key):
            q, q_src = int(QTY_IN_ITEM.match(key).group(1)), "Item"
        r.qty = q or 1

        # Unit price: Amount column; if it's empty, the money in "Quantity x Total".
        # "P2.7k x 10pcs" = unit x qty, but "P5,800 (2pcs.)" / "P84,000 = 30pcs." = a total.
        unit = money(c["amount"])
        qty_text = str(c["qty_total"] or "")
        each = re.match(r"\s*P?\s*([\d,.]+k?)\s*[xX]\s*\d+", qty_text)          # "4,999x2pcs ..." -> 4,999 each
        if unit is not None and r.qty > 1:
            if isinstance(c["amount"], str) and re.search(r"(\(|=)\s*\d+\s*pcs?", c["amount"], re.I):
                unit = unit / r.qty                                          # "P5,800 (2pcs.)" is a total
            elif q_src == "Item":
                unit = unit / r.qty                                          # "4 bottles ink" @ 400 = the lot
            elif each and money(each.group(1)) and \
                    abs(unit - money(each.group(1)) * r.qty) < abs(unit - money(each.group(1))):
                unit = unit / r.qty                                          # Amount 11000 = 2 x 4,999 + fees
        if unit is None:
            total = money(c["qty_total"])
            if total is not None:
                pcs_total = re.search(r"\d+\s*pcs?\s*-\s*P?\s*([\d,]+)", str(c["qty_total"] or ""), re.I)
                unit = total if not pcs_total else money(pcs_total.group(1)) / r.qty
        r.unit_price = unit

        plain = lambda v: v is None or isinstance(v, (int, float)) or \
            re.fullmatch(r"(1pc - )?P?\s*[\d,]+", str(v).strip()) is not None
        if not (plain(c["qty_total"]) and plain(c["amount"])) or r.qty != 1:
            r.flags.append(f"parsed qty={r.qty} (from {q_src if q else 'default'}), "
                           f"unit price={r.unit_price} from Amount={c['amount']!r}, Qty x Total={c['qty_total']!r}")
        stated_total = re.search(r"=\s*P?\s*([\d,]{4,})\s*$", qty_text)
        if stated_total and r.unit_price and r.qty == 1:
            ratio = money(stated_total.group(1)) / r.unit_price
            if abs(ratio - round(ratio)) > 0.05:
                r.flags.append(f"AMBIGUOUS: total {stated_total.group(1)} is {ratio:.1f}x the unit price, "
                               "not a whole number of units; imported as qty 1, check the sheet")

        r.paid = is_paid(c["payment_status"])
        text = " ".join(str(c[k] or "") for k in ("amount", "qty_total", "remarks")).lower()
        if r.paid and "not paid" in text:
            r.paid = False
            c = {**c, "payment_status": f"{c['payment_status']} (but marked 'not paid' in the row)"}
        if not r.paid:
            report["unpaid"].append((n, c, r))
            continue
        rows.append(r)

    # Price sanity check: a unit price far from the product's usual price is usually a
    # misread total (or a typo in the sheet).
    prices = defaultdict(list)
    for r in rows:
        if r.unit_price:
            prices[r.key].append(r.unit_price)
    for r in rows:
        ps = sorted(prices[r.key])
        if r.unit_price is None:
            r.unit_price = ps[len(ps) // 2] if ps else 0.0
            r.flags.append(f"no price in the row: used this product's median price {r.unit_price:,.0f}")
        if len(ps) >= 3 and r.unit_price:
            median = ps[len(ps) // 2]
            if not 0.5 <= r.unit_price / median <= 2:
                r.flags.append(f"unit price {r.unit_price:,.0f} is far from this product's median {median:,.0f}")
    report["review"] = [r for r in rows if r.flags]
    return rows, report


# --------------------------------------------------------------------------- report
def money_s(v):
    return f"P{v:,.0f}"


def print_report(rows: list[Row], report: dict, path: Path):
    def head(t):
        print(f"\n{'=' * 100}\n{t}\n{'=' * 100}")

    total_lines = len(report["separator"]) + len(report["no_item"]) + len(report["services"]) + \
        len(report["unpaid"]) + len(rows)
    head(f"SOURCE  {path.name}  (tab SALES, data from row 3)")
    print(f"  rows read (to the last used row):   {total_lines}")
    print(f"  blank separator rows skipped:       {len(report['separator'])}")
    print(f"  rows with data but no item skipped: {len(report['no_item'])}")
    print(f"  services excluded:                  {len(report['services'])}")
    print(f"  product rows not paid (skipped):    {len(report['unpaid'])}")
    print(f"  PAID product sales to import:       {len(rows)}")

    head("EXCLUDED: SERVICES (repair / reset / labor, not stock)")
    for n, c in report["services"]:
        print(f"  row {n:>3}  {str(c['date'].date()) if isinstance(c['date'], dt.datetime) else '':<10}  "
              f"{c['item'][:80]!r}  [{c['model']}]  {c['amount']}  status={c['payment_status']}")

    head("SKIPPED: rows with data but no Item (can't tell what was sold)")
    for n, c in report["no_item"]:
        print(f"  row {n:>3}  " + ", ".join(f"{k}={v!r}" for k, v in c.items() if v not in (None, "")))

    head("SKIPPED: product rows that are not paid")
    by_status = Counter(str(c["payment_status"]) for _, c, _ in report["unpaid"])
    for status, k in by_status.most_common():
        print(f"  {k:>3}  payment status = {status}")
    for n, c, r in report["unpaid"]:
        print(f"        row {n:>3} {r.date}  {r.item} {r.model or ''}  x{r.qty}  status={c['payment_status']}")

    head("ITEM NAME MAPPING (raw sheet text -> product item)")
    for item in sorted(report["item_map"]):
        print(f"  {item:<30} <- " + ", ".join(f"{raw!r} x{k}" for raw, k in report["item_map"][item].most_common()))

    head("MODEL MAPPING (raw -> clean), per item")
    for item in sorted(report["model_map"]):
        print(f"  {item}: " + ", ".join(f"{m} x{k}" for m, k in sorted(report["model_map"][item].items())))

    head("CLEANED PRODUCTS (paid sales only), most units first")
    by_product = defaultdict(list)
    for r in rows:
        by_product[r.key].append(r)
    print(f"  {'product':<44} {'brand':<8} {'category':<15} {'sales':>5} {'units':>5} {'revenue':>11}  first       last")
    for key, rs in sorted(by_product.items(), key=lambda kv: (-sum(r.qty for r in kv[1]), product_name(kv[0]))):
        print(f"  {product_name(key)[:44]:<44} {(rs[0].brand or '-'):<8} {rs[0].category:<15} {len(rs):>5} "
              f"{sum(r.qty for r in rs):>5} {money_s(sum((r.unit_price or 0) * r.qty for r in rs)):>11}  "
              f"{min(r.date for r in rs)}  {max(r.date for r in rs)}")
    print(f"\n  {len(by_product)} products, {len(rows)} sales, {sum(r.qty for r in rows)} units, "
          f"{money_s(sum((r.unit_price or 0) * r.qty for r in rows))} revenue")
    print(f"  date span of paid sales: {min(r.date for r in rows)} to {max(r.date for r in rows)}")

    head("NEEDS REVIEW: imported rows whose values were inferred or parsed from free text")
    for r in report["review"]:
        print(f"  row {r.sheet_row:>3} {r.date}  {product_name(r.key)}")
        for f in r.flags:
            print(f"           - {f}")

    pay = Counter(payment_method(r.raw["payment_method"])[1] for r in rows)
    head("KNOWN LIMITATIONS")
    print(f"  * stock_quantity is NOT real: every product is created with stock = {DEFAULT_STOCK} (the sheet has no stock counts).")
    print(f"  * all {len(rows)} orders belong to one placeholder customer '{IMPORT_USER}' (the sheet has no customer IDs).")
    print(f"  * payment method: {pay['GCASH']} GCASH, {pay['BANK_TRANSFER']} bank transfer, and {pay['fallback']} other "
          "(Cash/COD/Shopee/LBC/mixed/blank) stored as BANK_TRANSFER, since the schema ENUM has no value for them.")
    print("  * order time is unknown: every order is stamped 12:00 on its sheet date.")


def product_name(key) -> str:
    item, model = key
    return f"{item} {model}" if model else item


def payment_method(raw) -> tuple[str, str]:
    s = (raw or "").lower() if isinstance(raw, str) else ""
    if s.startswith("gcash") and "cash" not in s[5:]:
        return "GCASH", "GCASH"
    if "bdo" in s or "bank" in s:
        return "BANK_TRANSFER", "BANK_TRANSFER"
    return "BANK_TRANSFER", "fallback"


# --------------------------------------------------------------------------- import
def import_rows(rows: list[Row], database: str, s) -> dict:
    conn = pymysql.connect(host=s.db_host, port=s.db_port, user=s.db_user, password=s.db_password,
                           charset="utf8mb4", autocommit=False)
    with conn.cursor() as cur:
        cur.execute(f"DROP DATABASE IF EXISTS `{database}`")
        cur.execute(f"CREATE DATABASE `{database}` CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci")
        cur.execute(f"USE `{database}`")
    conn.autocommit(True)
    sqlscript.run_file(conn, SCHEMA)
    sqlscript.run_file(conn, MIGRATION)
    conn.autocommit(False)

    with conn.cursor() as cur:
        # Same categories as production, plus Printers (the sheet sells whole printers).
        cur.execute(f"SELECT COUNT(*) FROM information_schema.tables WHERE table_schema = %s AND table_name = 'categories'",
                    (PRODUCTION_DB,))
        have_prod = cur.fetchone()[0] == 1
        if have_prod:
            cur.execute(f"INSERT INTO categories (category_id, name) SELECT category_id, name FROM `{PRODUCTION_DB}`.categories")
            # Admin accounts are copied so you can log in to verify; customers are not.
            cur.execute(f"""INSERT INTO users (user_id, username, email, password_hash, role, created_at)
                            SELECT user_id, username, email, password_hash, role, created_at
                            FROM `{PRODUCTION_DB}`.users WHERE role = 'admin'""")
        admins_copied = cur.rowcount if have_prod else 0
        for name in ("Printer Parts", "Computer Parts", "Laptop Parts", "Accessories", "Electronics", "Other", "Printers"):
            cur.execute("INSERT IGNORE INTO categories (name) VALUES (%s)", (name,))
        cur.execute("SELECT name, category_id FROM categories")
        categories = dict(cur.fetchall())

        cur.execute("SELECT user_id FROM users WHERE role = 'admin' ORDER BY user_id LIMIT 1")
        admin = cur.fetchone()
        # Unusable password hash: the placeholder customer can never log in.
        cur.execute("INSERT INTO users (username, email, password_hash, role) VALUES (%s, %s, %s, 'user')",
                    (IMPORT_USER, "import_walkin@bongalink.invalid", "!import-placeholder-no-login"))
        customer = cur.lastrowid
        creator = admin[0] if admin else customer

        by_product = defaultdict(list)
        for r in rows:
            by_product[r.key].append(r)
        product_ids = {}
        for key, rs in sorted(by_product.items(), key=lambda kv: min(r.date for r in kv[1])):
            first = min(r.date for r in rs)
            last_price = max(rs, key=lambda r: (r.date, r.sheet_row)).unit_price or 0
            cur.execute(
                """INSERT INTO products (name, category_id, brand, price, stock_quantity, description, created_by, created_at)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s)""",
                (product_name(key), categories[rs[0].category], rs[0].brand, last_price, DEFAULT_STOCK,
                 f"Imported from the 2026 sales sheet. Stock count is a placeholder ({DEFAULT_STOCK}), not a real count.",
                 creator, dt.datetime.combine(first, dt.time(9, 0))),
            )
            product_ids[key] = cur.lastrowid

        orders = 0
        for r in sorted(rows, key=lambda r: (r.date, r.sheet_row)):
            unit = r.unit_price or 0
            method, _ = payment_method(r.raw["payment_method"])
            cur.execute(
                """INSERT INTO orders (user_id, contact_email, status, total_amount, payment_method, shipping_fee, created_at)
                   VALUES (%s, NULL, 'completed', %s, %s, 0, %s)""",
                (customer, unit * r.qty, method, dt.datetime.combine(r.date, dt.time(12, 0))),
            )
            cur.execute("INSERT INTO order_items (order_id, product_id, quantity, unit_price) VALUES (%s, %s, %s, %s)",
                        (cur.lastrowid, product_ids[r.key], r.qty, unit))
            orders += 1
    conn.commit()
    conn.close()
    return {"products": len(product_ids), "orders": orders, "admins_copied": admins_copied}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--file", type=Path, default=DEFAULT_FILE)
    ap.add_argument("--database", default="bongalink_real")
    ap.add_argument("--dry-run", action="store_true", help="clean and report only; write nothing")
    args = ap.parse_args()

    if args.database == PRODUCTION_DB:
        raise SystemExit(f"refusing to import into '{PRODUCTION_DB}' (production): this script drops and "
                         "recreates its target. Import into a copy and verify first.")
    rows, report = clean(args.file)
    print_report(rows, report, args.file)
    if args.dry_run:
        print("\nDRY RUN: nothing was written to any database.")
        return 0

    s = config.load()
    print(f"\nIMPORTING into database '{args.database}' on {s.db_host}:{s.db_port} "
          f"(dropped and recreated; production '{PRODUCTION_DB}' is only read for categories/admin accounts)")
    result = import_rows(rows, args.database, s)
    print(f"done: {result['products']} products, {result['orders']} orders (status 'completed') in '{args.database}'"
          + (f"; {result['admins_copied']} admin account(s) copied from '{PRODUCTION_DB}'" if result["admins_copied"] else
             f"\nNOTE: '{PRODUCTION_DB}' has no admin accounts to copy, so '{args.database}' has none. To log in to the "
             "dashboard, create one as in Backend/README.md (section 2), using this database name."))
    return 0


if __name__ == "__main__":
    sys.exit(main())
